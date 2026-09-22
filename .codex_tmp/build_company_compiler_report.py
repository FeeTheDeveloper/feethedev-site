from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path(r"C:\Users\uveav\OneDrive\Documents\GitHub\feethedev-site\outputs\FTD_OS_Company_Compiler_Full_Flow_and_Engineering_Report.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
MID_GRAY = "D9D9D9"
TEXT = RGBColor(32, 32, 32)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), MID_GRAY)


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def format_table(table, widths=None, header=True, center_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        set_col_widths(table, widths)
    if header and table.rows:
        set_repeat_table_header(table.rows[0])
    center_cols = set(center_cols or [])
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0 and header:
            for cell in row.cells:
                set_cell_shading(cell, NAVY)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
        elif r_idx % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                if c_idx in center_cols and r_idx > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
                    if r_idx > 0:
                        run.font.size = Pt(8.5)


def add_table(doc, headers, rows, widths, center_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    format_table(table, widths, center_cols=center_cols)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_number(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_label_para(doc, label, text):
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.13

for style_name, size, before, after in [
    ("Title", 22, 0, 18),
    ("Heading 1", 15, 12, 6),
    ("Heading 2", 12, 9, 4),
    ("Heading 3", 10.5, 7, 3),
]:
    st = styles[style_name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
    styles[list_style].font.name = "Arial"
    styles[list_style].font.size = Pt(10.25)
    styles[list_style].paragraph_format.space_after = Pt(3)

if "Flow Step" not in styles:
    fs = styles.add_style("Flow Step", WD_STYLE_TYPE.PARAGRAPH)
else:
    fs = styles["Flow Step"]
fs.font.name = "Arial"
fs.font.size = Pt(10.5)
fs.font.bold = True
fs.font.color.rgb = RGBColor(23, 54, 93)
fs.paragraph_format.left_indent = Inches(0.25)
fs.paragraph_format.space_before = Pt(3)
fs.paragraph_format.space_after = Pt(3)

title = doc.add_paragraph(style="Title")
title.add_run("FTD OS Company Compiler Full Flow and Engineering Report")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = subtitle.add_run("Sole owner operating model and multi tenant provisioning architecture")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(65, 65, 65)
meta = doc.add_paragraph()
meta.add_run("Prepared for Alfreddie Postell II\n").bold = True
meta.add_run("Engineering context date 22 September 2026\n")
meta.add_run("Status Architecture direction and implementation plan")

doc.add_heading("Executive decision", level=1)
doc.add_paragraph(
    "FTD_OS should use the Company Compiler as the single intake and planning system for every wholly owned company. "
    "The compiler may generate organizations with 1,000 or more logical positions, but it must separate organizational capacity from live provider accounts, paid licenses, service entitlements, and production authority."
)
doc.add_paragraph(
    "Alfreddie Postell II is the sole ultimate owner and final human authority for the companies and assets in scope. "
    "Shared ownership, minority approval, joint venture consent, and outside equity approval are disabled by default. "
    "Company, tenant, credential, data, budget, and audit boundaries remain separate because ownership does not merge provider accounts or security boundaries."
)
doc.add_paragraph(
    "The immediate engineering objective is a version controlled control plane that turns an approved COMPANY_BUILD_REQUEST into immutable plans, executes only tenant bound and human approved changes, verifies provider state, and preserves evidence and rollback data. "
    "The current workbooks and Apps Script files are planning sources and test fixtures. They are not production authorization."
)

doc.add_heading("Source plans and conclusions", level=1)
source_rows = [
    ("Hutchrok 1000 Agent Organization", "1,000 positions, 15 departments, 45 teams, 85 groups, 6,045 memberships", "Detailed organization and authority blueprint", "Planning only"),
    ("RunnerGang Provisioning Flow", "Spreadsheet driven external membership import", "Batching, retry, locking, resumability", "Unsafe credential and tenant boundary"),
    ("FTD Workspace Org Provisioner", "Four parents, twenty companies, standardized nested groups", "Declarative portfolio topology", "No tenant proof, approval artifact, or rollback"),
    ("Fee The Developer 1000 User Implementation", "1,000 staged accounts, 2,000 aliases, five batches, 3,000 memberships", "Useful generated execution package", "License and conflict evidence still missing"),
    ("Company Creation Intake", "Twenty four intake sections and separate provisioning approval", "Best candidate for canonical front door", "Must compile into validated schemas"),
]
add_table(doc, ["Source", "Scope", "Value", "Current classification"], source_rows, [1.65, 2.0, 2.0, 1.3])

doc.add_heading("Chronological engineering direction", level=1)
steps = [
    ("1", "Hutchrok organization model", "Established a complete 1,000 position hierarchy with departments, teams, OUs, groups, reporting routes, approval tiers, and staged activation waves."),
    ("2", "RunnerGang execution model", "Added resumable batch processing but exposed privileged execution to editable spreadsheet state and lacked authoritative tenant binding."),
    ("3", "FTD portfolio topology", "Defined repeatable parent, company, and department group creation, but placed desired state directly inside executable source."),
    ("4", "FTD 1,000 account package", "Expanded the blueprint into user import staging, alias plans, group memberships, batches, and an administrative runbook."),
    ("5", "Company Compiler intake", "Generalized company creation into a reusable request contract that distinguishes unknown facts, generated structure, simulation, and external provisioning."),
    ("6", "Sole owner decision", "Set one ultimate owner and final human authority while preserving independent company and provider boundaries."),
    ("7", "Current target", "Build a reusable OS side compiler and control plane with plan, approve, apply, verify, audit, and rollback stages."),
]
add_table(doc, ["Order", "Milestone", "Engineering meaning"], steps, [0.55, 1.9, 4.55], center_cols=[0])

doc.add_heading("Sole owner accessibility model", level=1)
doc.add_paragraph(
    "The portfolio registry should declare Alfreddie Postell II as the 100 percent ultimate owner of every listed company and asset. "
    "This removes unnecessary co owner gates, but it does not justify shared credentials, merged tenant storage, or silent cross company access."
)

ownership_rows = [
    ("Ultimate owner", "Alfreddie Postell II", "Global portfolio visibility and final human approval"),
    ("Ownership percentage", "100 percent", "No minority or partner approval workflow"),
    ("Shared assets", "Disabled by default", "Asset use across companies requires an explicit service relationship"),
    ("Agent self approval", "Prohibited", "No agent can approve its own restricted action"),
    ("Privileged administration", "Human only", "Delegation remains capability specific and revocable"),
    ("Owner override", "Available and audited", "Emergency action creates an immutable event and never removes evidence"),
]
add_table(doc, ["Control", "Default", "Effect"], ownership_rows, [1.55, 1.7, 3.75])

doc.add_heading("Ownership and filing treatment", level=2)
add_label_para(doc, "Disabled ownership flags. ", "Shared ownership, minority consent, joint venture consent, outside equity approval, partner asset approval, and ownership dispute flags default to false.")
add_label_para(doc, "Retained compliance flags. ", "Formation records, annual reports, taxes, beneficial ownership reports where applicable, licenses, permits, insurance, payroll, regulated activity, government contracting, DBAs, and intercompany accounting remain subject to verification.")
add_label_para(doc, "Evidence rule. ", "The owner's statement is authoritative for internal design. External filings, domains, provider tenants, licenses, and permissions remain NEEDS_VERIFICATION until supported by authoritative records.")

doc.add_heading("Accessibility hierarchy", level=2)
for line in [
    "Owner Control Plane",
    "  Company registry and global read visibility",
    "  Tenant access bindings",
    "    Company administrators",
    "    Human operators",
    "    Agent identities",
    "    Service identities",
    "  Approval and emergency suspension",
    "  Audit and resource metering",
]:
    p = doc.add_paragraph(style="Flow Step")
    p.add_run(line)

doc.add_heading("Core model distinctions", level=1)
model_rows = [
    ("Enterprise tenant", "Legal or operating entity that owns resources"),
    ("Workspace connection", "Verified binding to one provider tenant using immutable provider identifiers"),
    ("Company build request", "Submitted intake, uploaded evidence, requested outputs, and unresolved information"),
    ("Organization blueprint", "Departments, teams, jobs, reporting routes, and logical positions"),
    ("Logical position", "An organizational slot that consumes no license by itself"),
    ("Identity", "A human, agent, service, contractor, partner, or simulation identity assigned to a position"),
    ("Provider account", "A real account created in Google Workspace or another SaaS provider"),
    ("Capability", "A narrowly defined operation the control plane may authorize"),
    ("Entitlement", "A provider authorized license, feature, quota, credit, or consumption right"),
    ("Group", "A communication or routing container that does not independently grant operational authority"),
    ("Approval", "Human authorization tied to an immutable plan and expiration"),
    ("Audit event", "Actor, tenant, capability, approval, request, result, and affected resource"),
]
add_table(doc, ["Object", "Definition"], model_rows, [1.65, 5.35])

doc.add_heading("Capacity model", level=2)
doc.add_paragraph(
    "The schema must support five, 1,000, or substantially more positions without changing its architecture. "
    "The value 1,000 is configuration. The compiler must calculate the following quantities independently."
)
for item in [
    "Logical capacity and reserved positions",
    "Active positions and assigned identities",
    "Human, agent, service, contractor, partner, and simulation identities",
    "Provider accounts required",
    "Paid licenses required and available",
    "Product entitlements and quotas",
    "Storage, API, model, execution, and spending limits",
]:
    add_bullet(doc, item)

doc.add_heading("Complete Company Compiler flow", level=1)
flow_rows = [
    ("1", "Receive intake", "Human form, uploaded company package, internal blueprint, or approved prompt", "COMPANY_BUILD_REQUEST draft"),
    ("2", "Normalize input", "Validate field types, identifiers, ownership defaults, requested stage, and uploads", "Canonical request"),
    ("3", "Classify evidence", "Separate verified facts, owner assertions, generated structure, and missing external facts", "Evidence register and NEEDS_INFORMATION list"),
    ("4", "Research gaps", "Research only authorized missing facts and retain citations and dates", "Research record"),
    ("5", "Resolve company identity", "Assign company ID, legal and brand names, portfolio parent, ownership, and lifecycle state", "Company registry entry"),
    ("6", "Select tenant strategy", "Choose an existing tenant or propose a new isolated environment", "Tenant assignment proposal"),
    ("7", "Verify domain and provider", "Confirm customer ID, primary and secondary domains, license inventory, and administrative subject", "Workspace discovery snapshot"),
    ("8", "Generate organization", "Create departments, teams, jobs, reporting routes, and logical positions", "Organization plan"),
    ("9", "Generate identity demand", "Assign identity types and determine which positions require real provider accounts", "Identity plan"),
    ("10", "Assign capabilities", "Map jobs to capabilities, constraints, approval tiers, and separation of duties", "Authorization plan"),
    ("11", "Inventory entitlements", "Match identities and workloads to provider licenses, quotas, and service eligibility", "Entitlement plan"),
    ("12", "Generate resources", "Plan OUs, groups, aliases, mailboxes, service accounts, repositories, databases, and policies", "Resource plan"),
    ("13", "Simulate", "Test workflows, failures, approvals, capacity, and recovery without public activity", "Simulation report"),
    ("14", "Produce provisioning diff", "Compare desired state with live read only discovery using immutable IDs", "Provisioning plan"),
    ("15", "Approve", "Bind owner approval to tenant, plan hash, capability list, limits, and expiration", "Approval record"),
    ("16", "Apply", "Execute bounded, idempotent changes using a tenant specific connector", "Execution record"),
    ("17", "Verify", "Read provider state and reconcile every resource and membership", "Verification report"),
    ("18", "Activate", "Release only verified identities and capabilities for approved workloads", "Activation record"),
    ("19", "Operate", "Route work, meter consumption, monitor drift, and preserve audit events", "Operational ledger"),
    ("20", "Recover or retire", "Suspend, roll back eligible changes, revoke access, archive evidence, and release licenses", "Rollback or retirement record"),
]
add_table(doc, ["Step", "Stage", "Required action", "Output"], flow_rows, [0.45, 1.2, 3.6, 1.75], center_cols=[0])

doc.add_heading("Canonical machine artifacts", level=1)
artifact_rows = [
    ("company-build-request.json", "Intake selections, owner assertions, uploads, requested outputs, and missing information"),
    ("evidence-register.json", "Source, authority, verification state, date, and applicable company"),
    ("company-blueprint.json", "Company identity, business model, tenant strategy, technology, and portfolio relationships"),
    ("organization-plan.json", "Departments, teams, jobs, reporting routes, and logical positions"),
    ("identity-plan.json", "Identity types, account requirements, aliases, managers, and activation state"),
    ("authorization-plan.json", "Capabilities, approval tiers, constraints, and separation of duties"),
    ("entitlement-plan.json", "Licenses, quotas, eligibility rules, allocations, and meters"),
    ("workspace-discovery.json", "Read only provider state with immutable customer and resource IDs"),
    ("provisioning-plan.json", "Exact proposed changes, dependencies, preconditions, operation ceiling, and rollback eligibility"),
    ("approval-record.json", "Owner, plan hash, tenant, approved capabilities, limits, expiration, and signature"),
    ("execution-record.json", "Idempotency keys, API results, checkpoints, failures, and created or adopted resource IDs"),
    ("verification-report.json", "Expected and observed state with unresolved drift"),
    ("rollback-manifest.json", "Compensating actions and before state for eligible changes"),
]
add_table(doc, ["Artifact", "Purpose"], artifact_rows, [2.15, 4.85])
doc.add_paragraph(
    "Every artifact requires a schema version, company ID, tenant ID, creation time, source commit, content hash, lifecycle state, and predecessor references. "
    "Spreadsheets and DOCX reports are generated views of these artifacts rather than execution authority."
)

doc.add_heading("Authorization and capability design", level=1)
doc.add_paragraph(
    "The control plane should authorize named capabilities instead of granting an identity general control. "
    "Owner access remains broad, but every action is still tenant bound, logged, and subject to provider permissions."
)
cap_rows = [
    ("tenant.inspect", "Read customer, domain, license, OU, group, and policy state", "Read only"),
    ("organization.compile", "Generate organization and position plans", "Compiler"),
    ("deployment.upload", "Upload a reviewed script or configuration version", "Owner or deployment operator"),
    ("ou.create", "Create approved organizational units", "Approved plan"),
    ("group.create", "Create approved groups", "Approved plan"),
    ("group.settings.change", "Change visibility, posting, or membership policy", "Separate approval"),
    ("identity.create", "Create a provider user or service identity", "License proof and approval"),
    ("membership.internal.add", "Add an internal member", "Approved plan"),
    ("membership.external.add", "Add an external member", "Consent and separate approval"),
    ("connector.install", "Establish a provider connection", "Tenant administrator approval"),
    ("deployment.activate", "Enable a production execution path", "Owner approval"),
    ("execution.rollback", "Apply approved compensating actions", "Owner or incident approval"),
]
add_table(doc, ["Capability", "Meaning", "Minimum gate"], cap_rows, [1.7, 3.45, 1.85])

doc.add_heading("Cross company access", level=2)
doc.add_paragraph(
    "Common ownership permits centralized visibility but does not create implicit access between companies. "
    "Any company providing engineering, operations, marketing, finance, or another service to another company should receive a governed service relationship."
)
for item in [
    "Provider company and consumer company",
    "Authorized service and exact resources",
    "Permitted capabilities and prohibited actions",
    "Data classification and storage boundary",
    "Budget or spending ceiling",
    "Effective date, expiration, and revocation state",
    "Owner approval and audit reference",
]:
    add_bullet(doc, item)

doc.add_heading("OS build upload and execution architecture", level=1)
doc.add_paragraph(
    "The Windows workstation should host the developer interface, plan generator, validation tools, and reviewed deployment packages. "
    "It should not become the permanent vault for unrestricted multi tenant administrator credentials."
)

for line in [
    "Windows CLI and repository",
    "  Intake validator and schema registry",
    "  Compiler and generated artifacts",
    "  Read only discovery adapters",
    "  Plan and approval service",
    "  Tenant specific execution queue",
    "    Apps Script adapter for limited bootstrap work",
    "    Direct Admin API adapter for durable operations",
    "  Verification, audit, metering, and rollback",
]:
    p = doc.add_paragraph(style="Flow Step")
    p.add_run(line)

doc.add_heading("Local command interface", level=2)
commands = [
    ("inspect", "Read current tenant and provider state without mutation"),
    ("compile", "Generate canonical plans from an approved build request"),
    ("plan", "Produce an exact desired versus observed change set"),
    ("approve", "Create a signed and expiring approval for one plan hash"),
    ("upload", "Upload reviewed code or configuration without activating it"),
    ("apply", "Execute only the approved plan within operation limits"),
    ("verify", "Re read provider state and reconcile the execution"),
    ("status", "Report execution, drift, license, and entitlement state"),
    ("suspend", "Stop a tenant queue or capability"),
    ("rollback", "Apply eligible compensating actions from the recorded manifest"),
]
add_table(doc, ["Command", "Purpose"], commands, [1.25, 5.75])

doc.add_heading("Apps Script role", level=2)
doc.add_paragraph(
    "Apps Script may support tenant local bootstrap and limited administrative workflows. Source should be generated locally, reviewed in Git, uploaded as a version, and activated only through a separate approval. "
    "Editable spreadsheets must not control admin authorized background triggers. Long term multi tenant operations should use direct provider APIs through isolated connectors and a durable execution service."
)

doc.add_heading("Tenant expansion process", level=1)
tenant_steps = [
    "Register the company and proposed provider tenant without credentials.",
    "Verify the displayed tenant identity, immutable customer ID, and primary domain.",
    "Obtain read only discovery authorization from that tenant's administrator.",
    "Import current domains, licenses, users, groups, OUs, policies, and connector state.",
    "Compare discovery with the company blueprint and classify all differences.",
    "Review required scopes, administrative subject, and capability limits.",
    "Create an isolated tenant connection and secret binding.",
    "Generate a zero write plan and verify the operation ceiling.",
    "Obtain owner approval for a bounded pilot.",
    "Apply the pilot, verify it, and confirm rollback evidence.",
    "Expand in approved waves with independent checkpoints.",
    "Retain a tenant specific revocation and emergency suspension path.",
]
for item in tenant_steps:
    add_number(doc, item)

doc.add_heading("Status model", level=1)
status_rows = [
    ("DRAFT", "Input may change and has not passed validation"),
    ("NEEDS_INFORMATION", "A required factual input or authoritative record is missing"),
    ("VALIDATED", "Schema and internal consistency checks passed"),
    ("PROVISIONING_PLAN_READY", "A reviewed change plan exists but has no execution authority"),
    ("HUMAN_APPROVAL_REQUIRED", "The next step affects an external system or protected capability"),
    ("APPROVED_FOR_EXECUTION", "A current approval is bound to the exact tenant and plan hash"),
    ("EXECUTING", "Approved changes are in progress"),
    ("PARTIALLY_APPLIED", "Some changes succeeded and reconciliation is required"),
    ("VERIFIED", "Observed provider state matches the approved plan"),
    ("ROLLED_BACK", "Eligible changes were reversed and verified"),
    ("BLOCKED", "A required authority, provider state, or recovery condition prevents safe continuation"),
]
add_table(doc, ["State", "Meaning"], status_rows, [2.0, 5.0])

doc.add_heading("Current plan classifications", level=2)
current_rows = [
    ("Hutchrok 1,000 agent workbook", "PROVISIONING_PLAN_READY", "License, recovery, pilot, and audit prerequisites remain open"),
    ("FTD 1,000 user workbook", "PROVISIONING_PLAN_READY", "Dashboard wording is more optimistic than the unresolved gates"),
    ("RunnerGang script", "SECURITY REVIEW REQUIRED", "Mutable sheet data can steer privileged trigger execution"),
    ("FTD provisioner script", "SECURITY REVIEW REQUIRED", "No authoritative customer binding, approval record, or rollback"),
    ("Company Compiler intake", "VALIDATED DESIGN DIRECTION", "Machine schema and implementation remain to be built"),
]
add_table(doc, ["Asset", "Classification", "Reason"], current_rows, [2.0, 2.0, 3.0])

doc.add_heading("Material risks and required controls", level=1)
risk_rows = [
    ("CRITICAL", "Wrong tenant execution", "Bind every plan and credential to an immutable customer ID and fail closed"),
    ("CRITICAL", "Credential confused deputy", "Never let editable workbooks steer background admin credentials"),
    ("CRITICAL", "Cross tenant privilege spread", "Use isolated connectors, queues, secrets, logs, and revocation per tenant"),
    ("HIGH", "Unapproved mass account creation", "Require license evidence, exact counts, plan hash, expiration, and owner approval"),
    ("HIGH", "Existing resource adoption", "Verify immutable IDs, ownership, purpose, managers, and configuration before reuse"),
    ("HIGH", "Missing rollback", "Record before state and compensating actions before each mutation"),
    ("HIGH", "Group based approval bypass", "Evaluate capabilities and approval records independently of membership"),
    ("MEDIUM", "Quota and propagation failures", "Use structured retries, checkpoints, readiness checks, and bounded waves"),
    ("MEDIUM", "Mutable or incomplete audit", "Write append only external events before and after every operation"),
    ("MEDIUM", "Schema drift", "Version schemas and migrations and preserve artifact lineage"),
]
add_table(doc, ["Severity", "Risk", "Required control"], risk_rows, [0.9, 2.2, 3.9], center_cols=[0])

doc.add_heading("Implementation roadmap", level=1)
roadmap_rows = [
    ("Phase 0", "Freeze unsafe production entry points", "Retain read only discovery and dry run only", "No uncontrolled writes"),
    ("Phase 1", "Schema foundation", "Company, tenant, position, identity, capability, entitlement, approval, and audit schemas", "Versioned contracts"),
    ("Phase 2", "Compiler", "Convert intake into blueprint, organization, identity demand, and generated reports", "Deterministic builds"),
    ("Phase 3", "Discovery", "Read only Google Workspace and other provider adapters", "Authoritative observed state"),
    ("Phase 4", "Planning and approval", "Immutable diffs, policy checks, operation ceilings, and expiring approvals", "Execution ready plans"),
    ("Phase 5", "Single tenant pilot", "Group and OU creation for one verified tenant", "Verified apply and rollback"),
    ("Phase 6", "Identity provisioning", "Bounded user creation after license and recovery controls", "Controlled activation waves"),
    ("Phase 7", "Multi tenant operations", "Isolated connections, queues, metering, monitoring, and suspension", "Portfolio control plane"),
    ("Phase 8", "Provider expansion", "Additional SaaS entitlement and workload adapters", "Governed cross provider routing"),
]
add_table(doc, ["Phase", "Objective", "Scope", "Exit condition"], roadmap_rows, [0.75, 1.7, 3.25, 1.3], center_cols=[0])

doc.add_heading("Tests required before production", level=1)
test_sections = {
    "Tenant and identity": [
        "Wrong customer ID, domain, administrative subject, Cloud project, or deployment ID must fail closed.",
        "A manifest for one company cannot be paired with another tenant's credential.",
        "Revoked connections and changed administrator roles must stop new execution.",
    ],
    "Plans and approvals": [
        "Any change after planning invalidates the plan hash and approval.",
        "Expired, reused, broader, or unsigned approvals are rejected.",
        "Owner override is recorded and cannot erase prior evidence.",
    ],
    "Execution": [
        "Concurrent and repeated applies remain idempotent.",
        "Timeouts, quota limits, partial failures, and delayed provider propagation resume safely.",
        "Existing conflicting resources are not treated as harmless duplicates.",
    ],
    "Recovery": [
        "A partially completed wave can be reconciled and rolled back without removing preexisting resources.",
        "Audit storage failure stops mutation rather than producing an unaudited change.",
        "Emergency suspension prevents new operations without corrupting in progress evidence.",
    ],
    "Scale and entitlements": [
        "The compiler handles 5, 1,000, and larger configurations without schema changes.",
        "Logical positions do not consume provider licenses until an account is approved.",
        "Provider specific entitlement limits remain separate from identity counts.",
    ],
}
for heading, items in test_sections.items():
    doc.add_heading(heading, level=2)
    for item in items:
        add_bullet(doc, item)

doc.add_heading("Operational runbook", level=1)
runbook_rows = [
    ("1", "Inspect", "Confirm company, tenant, customer ID, domains, licensing, connector identity, and current resources", "Read only evidence"),
    ("2", "Compile", "Generate canonical artifacts from the approved intake and evidence register", "Versioned plan set"),
    ("3", "Review", "Resolve NEEDS_INFORMATION and all policy violations", "Validated inputs"),
    ("4", "Plan", "Produce exact create, update, adopt, and no change operations", "Immutable diff"),
    ("5", "Approve", "Owner approves the tenant, plan hash, capability set, limits, and expiration", "Approval record"),
    ("6", "Pilot", "Apply the smallest useful wave", "Pilot execution record"),
    ("7", "Verify", "Reconcile provider state by immutable ID", "Pilot verification"),
    ("8", "Expand", "Apply later waves only after the previous wave is verified", "Wave records"),
    ("9", "Activate", "Enable workloads and entitlements approved for production", "Activation record"),
    ("10", "Monitor", "Track drift, consumption, failures, and expiring approvals", "Operational ledger"),
    ("11", "Recover", "Suspend, reconcile, or roll back using the execution manifest", "Recovery evidence"),
]
add_table(doc, ["Order", "Action", "Requirement", "Evidence"], runbook_rows, [0.5, 1.0, 4.15, 1.35], center_cols=[0])

doc.add_heading("Engineering decisions to preserve", level=1)
for item in [
    "Alfreddie Postell II is the sole ultimate owner and final human authority for the portfolio in scope.",
    "All companies and assets remain separately addressable even when ownership is common.",
    "One owner receives unified visibility through explicit bindings rather than shared unrestricted credentials.",
    "Unknown external facts remain NEEDS_INFORMATION and are never invented.",
    "A logical position, identity, email address, alias, group, capability, and entitlement are separate objects.",
    "One thousand is a configurable capacity value rather than an architectural ceiling or automatic license count.",
    "Privileged administration remains human controlled.",
    "External execution and production activation require separate approval.",
    "Simulation activity remains separated from public customer activity.",
    "Every external mutation must be tenant bound, idempotent, verified, auditable, and recoverable.",
]:
    add_bullet(doc, item)

doc.add_heading("Recommended next engineering package", level=1)
doc.add_paragraph(
    "The next implementation package should create the repository structure, JSON schemas, sample COMPANY_BUILD_REQUEST, portfolio registry, read only tenant discovery interface, plan format, approval record, and a command line prototype supporting inspect, compile, plan, and verify. "
    "Apply, upload, connector installation, and production activation should remain disabled until the tenant binding, approval, audit, idempotency, and rollback tests pass."
)

doc.add_heading("Source inventory", level=1)
sources = [
    "Hutchrok_1000_Agent_Organization.xlsx dated 21 September 2026",
    "Fee_The_Developer_1000_User_Implementation.xlsx dated 22 September 2026",
    "RunnerGang_Provisioning_Flow.gs",
    "ftd-workspace-org-provisioner.gs",
    "Plays Ranch and FTD_OS Company Creation Intake Template",
    "Conversation decisions concerning authorized 1,000 seat capacity, sole ownership, accessibility, and multi tenant control",
]
for item in sources:
    add_bullet(doc, item)

doc.core_properties.title = "FTD OS Company Compiler Full Flow and Engineering Report"
doc.core_properties.subject = "Company Compiler architecture, ownership, accessibility, provisioning, and operations"
doc.core_properties.author = "Fee The Developer"
doc.core_properties.keywords = "FTD_OS, Company Compiler, provisioning, tenant isolation, identity, entitlement, audit"
doc.save(OUT)
print(str(OUT))
